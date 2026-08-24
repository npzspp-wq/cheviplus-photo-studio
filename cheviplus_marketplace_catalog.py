"""Marketplace catalog for Cheviplus Photo Studio.

Imports a large Excel applicability database into a local SQLite index once.
Employees then type an article/catalog number and instantly get the matched
product plus consolidated vehicle applicability. No Internet or AI is required.
"""
from __future__ import annotations

import json
import re
import sqlite3
import threading
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import app

APP_VERSION = "5.16"
APP_BUILD = "2026.08.24.01"
CATALOG_DIR = app.APP_DIR / "marketplace_data"
CATALOG_DB = CATALOG_DIR / "catalog.sqlite3"
CATALOG_META = CATALOG_DIR / "catalog_meta.json"

_ORIGINAL_BUILD = app.App._build
_ORIGINAL_SETTINGS_PAYLOAD = app.App.settings_payload
_ORIGINAL_LOAD_SETTINGS = app.App.load_settings
_ORIGINAL_RESET_SETTINGS = app.App.reset_settings


def normalize_key(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    return re.sub(r"[^0-9A-ZА-ЯЁ]", "", text)


def clean_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _header_map(row):
    result = {}
    aliases = {
        "nomenclature": {"номенклатура", "наименование", "товар", "название"},
        "code": {"код", "код номенклатуры"},
        "article": {"артикул", "каталожный номер", "номер", "номер детали"},
        "brand": {"марка", "бренд авто"},
        "model": {"модель"},
        "model_year": {"модельный год", "модельныйгод"},
        "comment": {"комментарий", "примечание"},
        "years": {"года", "годы", "период"},
    }
    normalized = [str(v or "").strip().lower().replace("ё", "е") for v in row]
    for field, names in aliases.items():
        for idx, value in enumerate(normalized):
            if value in {n.replace("ё", "е") for n in names}:
                result[field] = idx
                break
    return result


def _cell(row, mapping, field):
    idx = mapping.get(field)
    if idx is None or idx >= len(row):
        return ""
    return clean_text(row[idx])


def import_excel_to_sqlite(xlsx_path: Path, progress=None) -> dict:
    """Stream XLSX rows into SQLite. Designed for 100k+ rows without huge RAM use."""
    from openpyxl import load_workbook

    CATALOG_DIR.mkdir(parents=True, exist_ok=True)
    temp_db = CATALOG_DB.with_suffix(".tmp.sqlite3")
    if temp_db.exists():
        temp_db.unlink()

    conn = sqlite3.connect(temp_db)
    try:
        conn.execute("PRAGMA journal_mode=OFF")
        conn.execute("PRAGMA synchronous=OFF")
        conn.execute("PRAGMA temp_store=MEMORY")
        conn.execute(
            """CREATE TABLE product_rows (
                id INTEGER PRIMARY KEY,
                article_norm TEXT NOT NULL,
                article TEXT,
                code_norm TEXT,
                code TEXT,
                nomenclature TEXT,
                brand TEXT,
                model TEXT,
                model_year TEXT,
                comment TEXT,
                years TEXT
            )"""
        )

        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        total_rows = 0
        imported = 0
        skipped = 0
        sheets_used = 0

        for ws in wb.worksheets:
            iterator = ws.iter_rows(values_only=True)
            try:
                header = next(iterator)
            except StopIteration:
                continue
            mapping = _header_map(header)
            if "article" not in mapping and "code" not in mapping:
                continue
            sheets_used += 1
            batch = []
            for row in iterator:
                total_rows += 1
                article = _cell(row, mapping, "article")
                code = _cell(row, mapping, "code")
                article_norm = normalize_key(article)
                code_norm = normalize_key(code)
                if not article_norm and not code_norm:
                    skipped += 1
                    continue
                batch.append((
                    article_norm, article, code_norm, code,
                    _cell(row, mapping, "nomenclature"),
                    _cell(row, mapping, "brand"),
                    _cell(row, mapping, "model"),
                    _cell(row, mapping, "model_year"),
                    _cell(row, mapping, "comment"),
                    _cell(row, mapping, "years"),
                ))
                if len(batch) >= 2000:
                    conn.executemany(
                        "INSERT INTO product_rows(article_norm,article,code_norm,code,nomenclature,brand,model,model_year,comment,years) VALUES (?,?,?,?,?,?,?,?,?,?)",
                        batch,
                    )
                    imported += len(batch)
                    batch.clear()
                    if progress:
                        progress(imported)
            if batch:
                conn.executemany(
                    "INSERT INTO product_rows(article_norm,article,code_norm,code,nomenclature,brand,model,model_year,comment,years) VALUES (?,?,?,?,?,?,?,?,?,?)",
                    batch,
                )
                imported += len(batch)

        conn.execute("CREATE INDEX idx_article_norm ON product_rows(article_norm)")
        conn.execute("CREATE INDEX idx_code_norm ON product_rows(code_norm)")
        conn.execute("CREATE INDEX idx_nomenclature ON product_rows(nomenclature)")
        conn.commit()
        wb.close()
    finally:
        conn.close()

    if CATALOG_DB.exists():
        CATALOG_DB.unlink()
    temp_db.replace(CATALOG_DB)

    metadata = {
        "source_name": xlsx_path.name,
        "source_path": str(xlsx_path),
        "rows": imported,
        "skipped": skipped,
        "sheets_used": sheets_used,
    }
    CATALOG_META.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return metadata


def catalog_metadata() -> dict:
    if not CATALOG_META.exists() or not CATALOG_DB.exists():
        return {}
    try:
        return json.loads(CATALOG_META.read_text(encoding="utf-8"))
    except Exception:
        return {}


def lookup_product(query: str) -> dict | None:
    key = normalize_key(query)
    if not key or not CATALOG_DB.exists():
        return None
    conn = sqlite3.connect(CATALOG_DB)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            """SELECT article,code,nomenclature,brand,model,model_year,comment,years
               FROM product_rows
               WHERE article_norm=? OR code_norm=?
               ORDER BY nomenclature,brand,model,years,model_year""",
            (key, key),
        ).fetchall()
    finally:
        conn.close()
    if not rows:
        return None

    # Choose the most common non-empty name/article/code and consolidate applicability.
    def most_common(field):
        counts = {}
        for r in rows:
            v = clean_text(r[field])
            if v:
                counts[v] = counts.get(v, 0) + 1
        return max(counts, key=counts.get) if counts else ""

    applicability = []
    seen = set()
    for r in rows:
        parts = [clean_text(r["brand"]), clean_text(r["model"])]
        vehicle = " ".join(p for p in parts if p).strip()
        years = clean_text(r["years"]) or clean_text(r["model_year"])
        comment = clean_text(r["comment"])
        line = vehicle
        if years:
            line += (" — " if line else "") + years
        if comment and comment.lower() not in line.lower():
            line += (" • " if line else "") + comment
        line = line.strip()
        norm = line.casefold()
        if line and norm not in seen:
            seen.add(norm)
            applicability.append(line)

    return {
        "query": query,
        "article": most_common("article"),
        "code": most_common("code"),
        "name": most_common("nomenclature"),
        "applicability": applicability,
        "row_count": len(rows),
    }


def _find_box(root, text):
    for widget in root.winfo_children():
        try:
            if isinstance(widget, ttk.LabelFrame) and widget.cget("text") == text:
                return widget
        except Exception:
            pass
        found = _find_box(widget, text)
        if found is not None:
            return found
    return None


def _format_result(product: dict, max_lines=12) -> str:
    lines = []
    if product.get("name"):
        lines.append(product["name"])
    ids = []
    if product.get("article"):
        ids.append(f"Артикул: {product['article']}")
    if product.get("code"):
        ids.append(f"Код: {product['code']}")
    if ids:
        lines.append("   ".join(ids))
    apps = product.get("applicability") or []
    if apps:
        lines.append("")
        lines.append("Применяемость:")
        lines.extend(f"• {x}" for x in apps[:max_lines])
        if len(apps) > max_lines:
            lines.append(f"… ещё {len(apps) - max_lines}")
    lines.append("")
    lines.append(f"Строк в базе: {product.get('row_count', 0)}")
    return "\n".join(lines)


def build_marketplace_catalog(self):
    _ORIGINAL_BUILD(self)
    self.marketplace_article_var = tk.StringVar(master=self, value="")
    self.marketplace_catalog_status = tk.StringVar(master=self, value="База не загружена")
    self.marketplace_result_var = tk.StringVar(master=self, value="Введите артикул или код номенклатуры")
    self.marketplace_selected_product = None

    basic = _find_box(self, "2. Основные настройки")
    if basic is None:
        return
    parent = basic.master
    box = ttk.LabelFrame(parent, text="1А. Маркетплейсы — база товаров", padding=12)
    box.pack(fill="x", pady=(10, 0), before=basic)

    top = ttk.Frame(box)
    top.pack(fill="x")
    ttk.Button(top, text="Загрузить Excel базу", command=lambda: self._marketplace_import_excel()).pack(side="left")
    ttk.Label(top, textvariable=self.marketplace_catalog_status).pack(side="left", padx=10)

    search = ttk.Frame(box)
    search.pack(fill="x", pady=(8, 0))
    ttk.Label(search, text="Артикул / код:").pack(side="left")
    entry = ttk.Entry(search, textvariable=self.marketplace_article_var, width=28)
    entry.pack(side="left", padx=6)
    ttk.Button(search, text="НАЙТИ ТОВАР", command=lambda: self._marketplace_lookup()).pack(side="left")
    entry.bind("<Return>", lambda _e: self._marketplace_lookup())

    result = tk.Label(
        box,
        textvariable=self.marketplace_result_var,
        justify="left",
        anchor="w",
        bg="#ffffff",
        relief="solid",
        borderwidth=1,
        padx=10,
        pady=8,
        wraplength=720,
    )
    result.pack(fill="x", pady=(8, 0))

    meta = catalog_metadata()
    if meta:
        self.marketplace_catalog_status.set(f"База: {meta.get('source_name','Excel')} • {meta.get('rows',0):,} строк".replace(",", " "))


def _marketplace_import_excel(self):
    path = filedialog.askopenfilename(
        title="Выберите Excel с базой товаров",
        filetypes=[("Excel", "*.xlsx")],
    )
    if not path:
        return
    self.marketplace_catalog_status.set("Импорт базы…")
    self.status.set("Импорт Excel для маркетплейсов…")

    def progress(count):
        if count % 10000 == 0:
            self.ui_queue.put(("marketplace_import_progress", count))

    def worker():
        try:
            meta = import_excel_to_sqlite(Path(path), progress=progress)
            self.ui_queue.put(("marketplace_import_done", meta))
        except Exception as exc:
            self.ui_queue.put(("marketplace_import_error", str(exc)))

    threading.Thread(target=worker, daemon=True).start()


def _marketplace_lookup(self):
    query = self.marketplace_article_var.get().strip()
    if not query:
        messagebox.showwarning("Маркетплейсы", "Введите артикул или код номенклатуры.")
        return
    if not CATALOG_DB.exists():
        messagebox.showwarning("Маркетплейсы", "Сначала загрузите Excel базу товаров.")
        return
    product = lookup_product(query)
    if not product:
        self.marketplace_selected_product = None
        self.marketplace_result_var.set(f"Товар не найден: {query}")
        self.status.set("Товар в базе не найден")
        return
    self.marketplace_selected_product = product
    self.marketplace_result_var.set(_format_result(product))
    self.status.set(f"Найден товар: {product.get('article') or product.get('code') or query}")


def process_ui_queue_marketplace(self):
    # This method is intentionally not installed; instead we extend the existing
    # queue handler by polling our own lightweight queue events through after().
    pass


def _poll_marketplace_events(self):
    # Existing app consumes ui_queue itself. We use a dedicated queue to avoid
    # stealing its events.
    try:
        while True:
            event = self._marketplace_queue.get_nowait()
            kind = event[0]
            if kind == "progress":
                count = event[1]
                self.marketplace_catalog_status.set(f"Импортировано: {count:,} строк".replace(",", " "))
            elif kind == "done":
                meta = event[1]
                self.marketplace_catalog_status.set(f"База: {meta.get('source_name','Excel')} • {meta.get('rows',0):,} строк".replace(",", " "))
                self.status.set("База маркетплейсов загружена")
                messagebox.showinfo("Маркетплейсы", f"База загружена.\nСтрок: {meta.get('rows',0):,}".replace(",", " "))
            elif kind == "error":
                self.marketplace_catalog_status.set("Ошибка импорта")
                self.status.set("Ошибка импорта базы")
                messagebox.showerror("Маркетплейсы", event[1])
    except Exception:
        pass
    self.after(150, lambda: _poll_marketplace_events(self))


def _marketplace_import_excel_v2(self):
    import queue
    if not hasattr(self, "_marketplace_queue"):
        self._marketplace_queue = queue.Queue()
        self.after(150, lambda: _poll_marketplace_events(self))
    path = filedialog.askopenfilename(title="Выберите Excel с базой товаров", filetypes=[("Excel", "*.xlsx")])
    if not path:
        return
    self.marketplace_catalog_status.set("Импорт базы…")
    self.status.set("Импорт Excel для маркетплейсов…")

    def progress(count):
        self._marketplace_queue.put(("progress", count))

    def worker():
        try:
            meta = import_excel_to_sqlite(Path(path), progress=progress)
            self._marketplace_queue.put(("done", meta))
        except Exception as exc:
            self._marketplace_queue.put(("error", str(exc)))
    threading.Thread(target=worker, daemon=True).start()


def settings_payload_marketplace(self):
    data = _ORIGINAL_SETTINGS_PAYLOAD(self)
    data["marketplace_last_article"] = self.marketplace_article_var.get() if hasattr(self, "marketplace_article_var") else ""
    return data


def load_settings_marketplace(self):
    _ORIGINAL_LOAD_SETTINGS(self)
    if not app.SETTINGS_FILE.exists() or not hasattr(self, "marketplace_article_var"):
        return
    try:
        data = json.loads(app.SETTINGS_FILE.read_text(encoding="utf-8"))
        self.marketplace_article_var.set(data.get("marketplace_last_article", ""))
    except Exception:
        pass


def reset_settings_marketplace(self):
    _ORIGINAL_RESET_SETTINGS(self)
    if hasattr(self, "marketplace_article_var"):
        self.marketplace_article_var.set("")


app.App._build = build_marketplace_catalog
app.App._marketplace_import_excel = _marketplace_import_excel_v2
app.App._marketplace_lookup = _marketplace_lookup
app.App.settings_payload = settings_payload_marketplace
app.App.load_settings = load_settings_marketplace
app.App.reset_settings = reset_settings_marketplace
app.APP_VERSION = APP_VERSION
app.APP_BUILD = APP_BUILD
